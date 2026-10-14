"""Exportação cadastral: uma linha por participante, sem credenciais de acesso."""

from collections import defaultdict
from datetime import datetime, timezone
from io import BytesIO
import re
from zipfile import ZipFile, ZIP_DEFLATED
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import selectinload
from sqlmodel import select

from app.models.all_models import Usuario, Ingresso, SolicitacaoValidacao, SolicitacaoCancelamento, UserRole


COLUNAS = (
    "ID", "Nome", "CPF", "Data de nascimento", "E-mail", "Telefone",
    "Secretaria", "Sigla da secretaria", "Setor", "Vínculo", "Cadastro ativo",
    "Vínculo validado", "Resgate bloqueado", "Último acesso", "Foto (caminho)",
    "Status da homologação", "ID da solicitação de validação", "Validação solicitada em",
    "Analisado em", "Analisado por", "Motivo da rejeição",
    "Comunicação registrada", "Comunicação registrada em",
    "Status do ingresso", "ID do ingresso", "Lote", "ID do lote",
    "Ingresso emitido em", "Ingresso cancelado em", "Motivo do cancelamento do ingresso",
    "Situação no evento", "Última movimentação", "Data da última movimentação",
    "Status da solicitação de cancelamento", "Cancelamento solicitado em",
    "Cancelamento analisado em", "Motivo da solicitação de cancelamento",
    "Total de ingressos no histórico", "Total de validações no histórico",
    "Lote do ingresso", "ID do lote do ingresso", "Lote da validação", "ID do lote da validação",
    "Cancelamento analisado por", "Motivo da rejeição do cancelamento", "Ingresso cancelado por",
)
EQUIPE = {role.value for role in UserRole if role != UserRole.servidor}
FUSO = ZoneInfo("America/Sao_Paulo")


def data_excel(value):
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(FUSO).replace(tzinfo=None)


def nascimento_excel(value):
    if not value:
        return None
    for formato in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, formato).date()
        except ValueError:
            pass
    return value


def mais_recente(itens, campo):
    return max(itens, key=lambda item: (
        data_excel(getattr(item, campo)) or datetime.min, item.id or 0
    ), default=None)


async def coletar_participantes(session):
    usuarios = (await session.execute(
        select(Usuario).where(Usuario.deleted_at.is_(None)).order_by(Usuario.nome, Usuario.id)
        .options(
            selectinload(Usuario.secretaria),
            selectinload(Usuario.ingressos).selectinload(Ingresso.lote),
            selectinload(Usuario.ingressos).selectinload(Ingresso.cancelado_por),
            selectinload(Usuario.solicitacoes_validacao).selectinload(SolicitacaoValidacao.lote),
            selectinload(Usuario.solicitacoes_validacao).selectinload(SolicitacaoValidacao.analisado_por),
            selectinload(Usuario.solicitacoes_cancelamento).selectinload(SolicitacaoCancelamento.analisado_por),
            selectinload(Usuario.logs_acesso),
        )
    )).scalars().all()
    registros = []
    for usuario in usuarios:
        # Considere tanto o papel principal quanto todos os papéis adicionais.
        papel = usuario.role.value if hasattr(usuario.role, "value") else str(usuario.role)
        if papel in EQUIPE or EQUIPE.intersection(usuario.get_roles_list()):
            continue
        validacao = mais_recente(usuario.solicitacoes_validacao, "criado_em")
        cancelamento = mais_recente(usuario.solicitacoes_cancelamento, "criado_em")
        ativo = mais_recente([i for i in usuario.ingressos if i.deleted_at is None], "data_resgate")
        ingresso = ativo or mais_recente(usuario.ingressos, "deleted_at")
        # O lote atual vem do ingresso ativo; sem ele, da última validação, ou do histórico.
        lote = ativo.lote if ativo else (validacao.lote if validacao else (ingresso.lote if ingresso else None))
        secretaria = usuario.secretaria
        log = mais_recente(usuario.logs_acesso, "data_hora")
        movimento = log.tipo.value if log and hasattr(log.tipo, "value") else (str(log.tipo) if log else None)
        status_validacao = {"pendente_analise": "Em validação", "aprovado": "Homologado", "rejeitado": "Rejeitado"}.get(
            validacao.status, validacao.status
        ) if validacao else ("Vínculo validado (sem solicitação)" if usuario.validado else "Sem solicitação")
        status_ingresso = "Ativo" if ativo else (
            "Extraviado / cancelado administrativamente" if ingresso and ingresso.cancelado_por_id is not None
            else "Cancelado" if ingresso else "Não emitido"
        )
        if ativo:
            situacao = "Presente no evento" if movimento == "entrada" else "Fora do evento"
        elif ingresso and ingresso.cancelado_por_id is not None:
            situacao = "Ingresso extraviado"
        elif ingresso or (cancelamento and cancelamento.status == "aprovado"):
            situacao = "Ingresso cancelado"
        elif validacao and validacao.status == "pendente_analise":
            situacao = "Em validação"
        elif validacao and validacao.status == "rejeitado":
            situacao = "Validação recusada"
        else:
            situacao = "Sem ingresso"
        dados = [
            usuario.id, usuario.nome, usuario.cpf, nascimento_excel(usuario.data_nascimento),
            usuario.email, usuario.telefone, secretaria.nome if secretaria else "Sem secretaria",
            secretaria.sigla if secretaria else None, usuario.setor, usuario.vinculo,
            "Sim" if usuario.ativo else "Não", "Sim" if usuario.validado else "Não",
            "Sim" if usuario.bloqueado_resgate else "Não", data_excel(usuario.ultimo_acesso), usuario.foto_rosto_url,
            status_validacao, validacao.id if validacao else None,
            data_excel(validacao.criado_em) if validacao else None,
            data_excel(validacao.analisado_em) if validacao else None,
            validacao.analisado_por.nome if validacao and validacao.analisado_por else None,
            validacao.motivo_rejeicao if validacao else None,
            ("Sim" if validacao.comunicado else "Não") if validacao else "Sem solicitação",
            data_excel(validacao.comunicado_em) if validacao else None,
            status_ingresso, ingresso.id if ingresso else None,
            lote.nome if lote else "Sem lote", lote.id if lote else None,
            data_excel(ingresso.data_resgate) if ingresso else None,
            data_excel(ingresso.deleted_at) if ingresso else None,
            ingresso.motivo_cancelamento if ingresso else None,
            situacao, {"entrada": "Entrada", "saida": "Saída"}.get(movimento, movimento),
            data_excel(log.data_hora) if log else None,
            {"pendente_codigo": "Aguardando código", "pendente_foto": "Aguardando foto",
             "pendente_analise": "Em análise", "aprovado": "Aprovado", "rejeitado": "Rejeitado"}.get(
                 cancelamento.status, cancelamento.status
             ) if cancelamento else "Sem solicitação",
            data_excel(cancelamento.criado_em) if cancelamento else None,
            data_excel(cancelamento.analisado_em) if cancelamento else None,
            cancelamento.motivo if cancelamento else None,
            len(usuario.ingressos), len(usuario.solicitacoes_validacao),
            ingresso.lote.nome if ingresso and ingresso.lote else None,
            ingresso.lote_id if ingresso else None,
            validacao.lote.nome if validacao and validacao.lote else None,
            validacao.lote_id if validacao else None,
            cancelamento.analisado_por.nome if cancelamento and cancelamento.analisado_por else None,
            cancelamento.motivo_rejeicao if cancelamento else None,
            ingresso.cancelado_por.nome if ingresso and ingresso.cancelado_por else None,
        ]
        registros.append({
            "dados": dados,
            "lote": (lote.id, lote.nome) if lote else (None, "Sem lote"),
            "secretaria": (secretaria.id, secretaria.nome) if secretaria else (None, "Sem secretaria"),
        })
    return registros


def nome_grupo(identifier, nome):
    # O ID distingue nomes iguais e continua visível mesmo após truncar o nome.
    seguro = re.sub(r'[\\/*?:\[\]<>|"\x00-\x1f]', "-", nome).strip(" .'") or "Grupo"
    return f"{identifier} - {seguro}" if identifier is not None else seguro


def criar_xlsx(grupos):
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Participantes — Serv indoor"
    workbook.properties.description = "Datas e horários de São Paulo. Comunicação registrada não confirma entrega SMTP."
    for nome, registros in grupos:
        sheet = workbook.create_sheet(nome[:31])
        sheet.append(COLUNAS)
        for row_number, registro in enumerate(registros, 2):
            for column_number, valor in enumerate(registro["dados"], 1):
                cell = sheet.cell(row=row_number, column=column_number,
                                  value=ILLEGAL_CHARACTERS_RE.sub("", valor) if isinstance(valor, str) else valor)
                if isinstance(cell.value, str):
                    # CPF, telefone e conteúdo iniciado por '=' permanecem texto literal.
                    cell.data_type = "s"
                    cell.number_format = "@"
                elif isinstance(cell.value, datetime):
                    cell.number_format = "dd/mm/yyyy hh:mm:ss"
                elif cell.is_date:
                    cell.number_format = "dd/mm/yyyy"
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 42
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="164E63")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for index, coluna in enumerate(COLUNAS, 1):
            sheet.column_dimensions[get_column_letter(index)].width = (
                10 if coluna in ("ID", "ID do lote", "ID do ingresso") else
                40 if coluna in ("Nome", "E-mail", "Secretaria", "Foto (caminho)") or "Motivo" in coluna else 25
            )
        sheet.sheet_view.showGridLines = True
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def gerar_exportacao(registros, modo, agrupar):
    if modo == "unico":
        return criar_xlsx([("Participantes", registros)]), "participantes.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    grupos = defaultdict(list)
    for registro in registros:
        grupos[registro[agrupar]].append(registro)
    organizados = [(nome_grupo(identifier, nome), itens)
                   for (identifier, nome), itens in sorted(grupos.items(), key=lambda grupo: (grupo[0][1].casefold(), grupo[0][0] or 0))]
    if not organizados:
        organizados = [("Participantes", [])]
    if modo == "abas":
        return criar_xlsx(organizados), f"participantes_por_{agrupar}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as pacote:
        for nome, itens in organizados:
            pacote.writestr(f"{nome[:100]}.xlsx", criar_xlsx([(nome, itens)]))
    return output.getvalue(), f"participantes_por_{agrupar}.zip", "application/zip"

import unittest
from datetime import datetime, timezone
from io import BytesIO
from zipfile import ZipFile

from fastapi import FastAPI
from openpyxl import load_workbook
from sqlmodel import SQLModel, Session, create_engine

from app.core.deps import get_current_user
from app.db.session import get_session
from app.models.all_models import Usuario, UserRole, Lote, Secretaria, Ingresso, SolicitacaoValidacao, SolicitacaoCancelamento, LogAcesso
from app.routers.admin import router
from app.services.participantes_export import coletar_participantes, gerar_exportacao, COLUNAS


class AsyncSession:
    def __init__(self, session):
        self.session = session

    async def execute(self, statement):
        return self.session.execute(statement)


class TestParticipantesExport(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        SQLModel.metadata.create_all(self.engine)
        self.db = Session(self.engine)
        self.session = AsyncSession(self.db)
        self.agora = datetime(2026, 10, 14, 15, 30, tzinfo=timezone.utc)
        self.db.add_all([
            Secretaria(id=1, nome="Secretaria / de Educação com um nome muito longo", sigla="EDU"),
            Secretaria(id=2, nome="Secretaria / de Educação com um nome muito longo", sigla="EDU2"),
            Lote(id=1, nome="Lote / com nome muito longo repetido para testar abas"),
            Lote(id=2, nome="Lote / com nome muito longo repetido para testar abas"),
        ])
        self.db.commit()
        for identifier in range(1, 71):
            self.db.add(Usuario(id=identifier, nome=f"Pessoa {identifier:03}", cpf=str(identifier).zfill(11),
                                secretaria_id=1 if identifier <= 35 else 2,
                                senha_hash="NAO_EXPORTAR_SENHA", session_id="NAO_EXPORTAR_SESSAO"))
        self.db.add_all([
            Usuario(id=80, nome="Sem secretaria e sem ingresso", cpf="00000000080", ativo=False),
            Usuario(id=81, nome="Excluído", cpf="00000000081", deleted_at=self.agora),
            Usuario(id=90, nome="Equipe principal", cpf="00000000090", role=UserRole.admin, roles="servidor"),
            Usuario(id=91, nome="Equipe adicional", cpf="00000000091", role=UserRole.servidor, roles="servidor,validador"),
        ])
        self.db.commit()
        self.db.add_all([
            SolicitacaoValidacao(id=1, usuario_id=1, lote_id=1, status="aprovado", comunicado=True,
                                analisado_em=self.agora, comunicado_em=self.agora, analisado_por_id=90),
            SolicitacaoValidacao(id=2, usuario_id=2, lote_id=2, status="pendente_analise"),
            SolicitacaoValidacao(id=3, usuario_id=3, lote_id=1, status="rejeitado", motivo_rejeicao="Vínculo não confirmado"),
            Ingresso(id=1, usuario_id=1, lote_id=1, qr_code_token="NAO_EXPORTAR_QR", data_resgate=self.agora),
            Ingresso(id=4, usuario_id=4, lote_id=2, qr_code_token="antigo", deleted_at=self.agora, cancelado_por_id=90),
            Ingresso(id=5, usuario_id=5, lote_id=2, qr_code_token="cancelado", deleted_at=self.agora),
            Ingresso(id=6, usuario_id=6, lote_id=2, qr_code_token="fora"),
            Ingresso(id=90, usuario_id=90, lote_id=1, qr_code_token="equipe"),
            LogAcesso(id=1, usuario_id=1, ingresso_id=1, data_hora=self.agora),
        ])
        self.db.add(SolicitacaoCancelamento(id=5, usuario_id=5, ingresso_id=5, token="cancelamento",
                                           codigo_verificacao="000000", codigo_expira_em=self.agora,
                                           status="aprovado", motivo="Desistência", analisado_em=self.agora,
                                           analisado_por_id=90))
        self.db.get(Usuario, 1).nome = '=HYPERLINK("https://example.invalid")'
        self.db.get(Usuario, 1).telefone = "00123456789"
        self.db.get(Usuario, 1).data_nascimento = "1990-01-02"
        self.db.commit()
        self.db.expunge_all()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    async def test_exporta_todos_sem_equipe_e_preserva_status_tipos_e_identificadores(self):
        registros = await coletar_participantes(self.session)
        conteudo, nome, media = gerar_exportacao(registros, "unico", "secretaria")
        self.assertEqual(nome, "participantes.xlsx")
        self.assertIn("spreadsheetml", media)
        workbook = load_workbook(BytesIO(conteudo))
        self.assertEqual(workbook.sheetnames, ["Participantes"])
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 72)
        linhas = {row[0]: dict(zip(COLUNAS, row)) for row in sheet.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(set(linhas), set(range(1, 71)) | {80})
        self.assertEqual(linhas[1]["CPF"], "00000000001")
        self.assertEqual(linhas[1]["Telefone"], "00123456789")
        self.assertEqual(linhas[1]["Status da homologação"], "Homologado")
        self.assertEqual(linhas[1]["Comunicação registrada"], "Sim")
        self.assertEqual(linhas[1]["Status do ingresso"], "Ativo")
        self.assertEqual(linhas[1]["Situação no evento"], "Presente no evento")
        self.assertEqual(linhas[1]["Analisado em"], datetime(2026, 10, 14, 12, 30))
        self.assertIsInstance(linhas[1]["Data de nascimento"], datetime)
        self.assertEqual(linhas[2]["Status da homologação"], "Em validação")
        self.assertEqual(linhas[2]["Status do ingresso"], "Não emitido")
        self.assertEqual(linhas[3]["Status da homologação"], "Rejeitado")
        self.assertEqual(linhas[4]["Status do ingresso"], "Extraviado / cancelado administrativamente")
        self.assertEqual(linhas[5]["Status do ingresso"], "Cancelado")
        self.assertEqual(linhas[5]["Status da solicitação de cancelamento"], "Aprovado")
        self.assertEqual(linhas[5]["Motivo da solicitação de cancelamento"], "Desistência")
        self.assertEqual(linhas[5]["Cancelamento analisado por"], "Equipe principal")
        self.assertEqual(linhas[6]["Situação no evento"], "Fora do evento")
        self.assertEqual(linhas[80]["Cadastro ativo"], "Não")
        self.assertEqual(sheet.freeze_panes, "C2")
        self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
        self.assertEqual(sheet.cell(2, 2).data_type, "s")
        self.assertFalse(any(cell.data_type == "f" for row in sheet for cell in row))
        self.assertNotIn("NAO_EXPORTAR", str(list(sheet.values)))

    async def test_abas_e_arquivos_por_lote_ou_secretaria_sem_perdas_ou_duplicacoes(self):
        registros = await coletar_participantes(self.session)
        for agrupar in ("lote", "secretaria"):
            for modo in ("abas", "arquivos"):
                with self.subTest(agrupar=agrupar, modo=modo):
                    conteudo, _, _ = gerar_exportacao(registros, modo, agrupar)
                    if modo == "arquivos":
                        with ZipFile(BytesIO(conteudo)) as pacote:
                            self.assertEqual(len(pacote.namelist()), 3)
                            workbooks = [load_workbook(BytesIO(pacote.read(name))) for name in pacote.namelist()]
                            self.assertTrue(all(len(w.sheetnames) == 1 for w in workbooks))
                    else:
                        workbooks = [load_workbook(BytesIO(conteudo))]
                        self.assertEqual(len(workbooks[0].sheetnames), 3)
                    ids = [row[0] for workbook in workbooks for sheet in workbook for row in sheet.iter_rows(min_row=2, values_only=True)]
                    self.assertEqual(len(ids), 71)
                    self.assertEqual(set(ids), set(range(1, 71)) | {80})
                    self.assertTrue(all(len(sheet.title) <= 31 for workbook in workbooks for sheet in workbook))

    async def test_exportacao_vazia_continua_sendo_arquivo_valido(self):
        for modo in ("unico", "abas", "arquivos"):
            conteudo, _, _ = gerar_exportacao([], modo, "lote")
            if modo == "arquivos":
                with ZipFile(BytesIO(conteudo)) as pacote:
                    conteudo = pacote.read(pacote.namelist()[0])
            workbook = load_workbook(BytesIO(conteudo))
            self.assertEqual(workbook.active.max_row, 1)

    async def test_endpoint_restringe_acesso_e_valida_opcoes(self):
        app = FastAPI()
        app.include_router(router)
        role = UserRole.servidor

        async def user():
            return Usuario(id=1000, nome="Operador", cpf="11144477735", role=role)

        async def session():
            yield self.session

        app.dependency_overrides[get_current_user] = user
        app.dependency_overrides[get_session] = session
        for role, query, esperado in [
            (UserRole.servidor, b"", 403), (UserRole.validador, b"", 403),
            (UserRole.admin, b"modo=invalido", 422), (UserRole.admin, b"agrupar=invalido", 422),
            (UserRole.admin, b"", 200), (UserRole.admin_geral, b"modo=arquivos&agrupar=lote", 200),
        ]:
            messages = []

            async def receive():
                return {"type": "http.request", "body": b"", "more_body": False}

            async def send(message):
                messages.append(message)

            await app({"type": "http", "asgi": {"version": "3.0"}, "method": "GET", "scheme": "http",
                       "path": "/api/admin/participantes/exportar-planilha", "query_string": query, "headers": []}, receive, send)
            start = next(m for m in messages if m["type"] == "http.response.start")
            self.assertEqual(start["status"], esperado)
            if esperado == 200:
                headers = dict(start["headers"])
                self.assertEqual(headers[b"cache-control"], b"no-store")
                self.assertIn(b"attachment", headers[b"content-disposition"])
